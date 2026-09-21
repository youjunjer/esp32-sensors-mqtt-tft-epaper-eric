# -*- coding: utf-8 -*-
"""Build the UTF-8 0914 IoT project report."""
from pathlib import Path
import io
from docx import Document
from PIL import Image
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\esp32實習")
OUT = ROOT / "0914物聯網成果報告.docx"
PHOTO = ROOT / "成果照片"
ARCH = ROOT / "架構圖_0914_imagegen.png"
FLOW = ROOT / "流程圖_0914_imagegen.png"
OLD_ARCH = ROOT / "系統架構圖.png"
OLD_FLOW = ROOT / "資料流程圖.png"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_text(cell, text, bold=False, color="17365D"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Microsoft JhengHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(color)

def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft JhengHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def add_para(doc, text="", style=None, size=10.5, bold=False, color="222222", align=None, before=0, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if align is not None: p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size, bold, color)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, 10)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, 15 if level == 1 else 12, True, "17365D" if level == 1 else "2F5597")
    return p

def add_picture(doc, path, width, caption=None):
    if not path.exists(): return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Normalize camera screenshots with missing/zero DPI metadata so Word can scale them.
    with Image.open(path) as im:
        buf = io.BytesIO()
        fmt = "PNG" if path.suffix.lower() == ".png" else "JPEG"
        if fmt == "JPEG" and im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        im.save(buf, format=fmt, dpi=(96, 96))
        buf.seek(0)
        p.add_run().add_picture(buf, width=Inches(width))
    if caption: add_para(doc, caption, size=9, color="666666", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True, "FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "17365D")
        table.rows[0].cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(table.rows) % 2 == 0: set_cell_shading(cells[i], "EAF2F8")
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("0914 物聯網成果報告  |  ")
    set_run_font(run, 8, color="777777")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.6)
sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
add_page_number(sec)
styles = doc.styles
styles["Normal"].font.name = "Microsoft JhengHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")

# Cover
add_para(doc, "0914", size=34, bold=True, color="17365D", align=WD_ALIGN_PARAGRAPH.CENTER, before=45, after=0)
add_para(doc, "物聯網成果報告", size=27, bold=True, color="2F5597", align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
add_para(doc, "0810 舊成果 × 0914 最新成果\nESP32 智慧環境監測與 MQTT 控制系統", size=13, color="555555", align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
add_picture(doc, ARCH, 6.25, "圖 1　以 ESP32 為核心的最新系統架構（imagegen 繪製）")
add_para(doc, "報告版本：2026 年 09 月 14 日", size=10, color="666666", align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
add_para(doc, "主程式：35_ili_mqtt_ctrl_page", size=10, color="666666", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# Preserve the original 0810 report as the first project stage.
add_heading(doc, "第一階段成果：0810 環境監測、Google Sheets 與 LINE", 1)
add_para(doc, "本節保留 0810 舊成果報告的主要內容，作為本專案由單向雲端紀錄逐步演進到 MQTT 雙向控制的基礎。")
add_heading(doc, "舊版摘要", 2)
add_para(doc, "第一階段以 ESP32 Dev Module 為核心，整合 DHT11 溫濕度感測器、光敏電阻與 I²C OLED 顯示器，完成環境資料的即時量測、視覺化呈現、Google Sheets 雲端紀錄與 LINE 異常通知。系統每秒更新感測值，每 10 秒將溫度、濕度及亮度資料寫入 Google Sheets；當溫度超過 28°C 或濕度超過 70% 時，透過 LINE Messaging API 通知指定使用者，異常期間每 30 秒最多發送一次。")
add_heading(doc, "舊版專案目標", 2)
add_bullets(doc, [
    "學習 ESP32 GPIO、ADC、I²C 與 Wi-Fi 網路功能。",
    "完成溫度、濕度與環境亮度的即時監測。",
    "以 OLED 提供現場即時資訊與網路傳輸狀態。",
    "將感測資料寫入 Google Sheets，方便保存與製作圖表。",
    "建立環境異常判斷，透過 LINE 即時通知使用者。",
])
add_heading(doc, "舊版系統架構與流程", 2)
add_para(doc, "舊版系統由感測端、ESP32 控制端、OLED 顯示端及 Google Sheets、LINE 兩個雲端服務組成。ESP32 負責讀取資料、判斷異常、更新畫面並透過 Wi-Fi 傳送資料。")
add_picture(doc, OLD_ARCH, 5.9, "圖 A　0810 舊版系統架構圖")
add_picture(doc, OLD_FLOW, 5.9, "圖 B　0810 舊版資料流程圖")
add_heading(doc, "舊版硬體、功能與測試紀錄", 2)
add_table(doc, ["元件", "用途", "ESP32 腳位", "備註"], [
    ["ESP32 Dev Module", "主控制器、Wi-Fi", "—", "Arduino ESP32 core 3.3.10"],
    ["DHT11", "溫度、濕度", "GPIO14", "每秒讀取一次"],
    ["光敏電阻", "環境亮度", "GPIO33 / ADC", "換算為 0–100%"],
    ["SSD1306 OLED", "即時顯示", "I²C", "顯示溫度、濕度、亮度"],
])
add_table(doc, ["功能", "舊版實作內容"], [
    ["即時監測", "DHT11 讀取溫度與濕度，GPIO33 ADC 讀取光敏電阻，亮度以 map() 轉換成 0–100%。"],
    ["OLED 顯示", "上方顯示溫度，下方左右顯示濕度與亮度，並使用溫度計、水滴、太陽圖示。"],
    ["Google Sheets", "以 HTTPS 連線 Google Apps Script，每 10 秒寫入溫度、濕度與亮度。"],
    ["LINE 通知", "溫度超過 28°C 或濕度超過 70% 時發送異常通知，異常期間限制發送頻率。"],
])
add_table(doc, ["測試項目", "預期結果", "結果"], [
    ["OLED 顯示", "持續顯示溫度、濕度、亮度", "通過"],
    ["Google Sheets", "每 10 秒新增一筆資料", "通過；可看到時間戳記與三項數值"],
    ["溫度異常", "溫度 > 28°C 時通知 LINE", "已完成通知流程"],
])
add_heading(doc, "舊版學習成果與階段結論", 2)
add_para(doc, "第一階段從單顆 LED、按鈕與類比輸入等基礎練習，逐步整合到感測器、OLED、Wi-Fi、Google Sheets 與 LINE Messaging API，實際練習 GPIO、ADC、I²C 顯示、HTTPS 傳輸、JSON／HTTP、計時器與異常通知邏輯。這些經驗成為後續改用 TFT 顯示、MQTT 雙向通訊及遠端 LED 控制的基礎。")
add_para(doc, "舊版主要檔案：25_dht_line/25_dht_line.ino；主要函式庫：WiFi、WiFiClientSecure、SimpleDHT、U8g2lib。")
doc.add_page_break()

add_heading(doc, "第二階段成果：0914 ILI9225 與 MQTT 雙向物聯網系統", 1)

add_heading(doc, "摘要", 1)
add_para(doc, "本專案以 ESP32 建置可連網的環境監測與控制裝置，讀取 DHT11 的溫度、濕度及光敏元件的亮度百分比，使用 ILI9225 176×220 TFT 顯示器即時呈現資料。系統透過 Wi-Fi 連線至 MQTT broker，每 10 秒以 JSON 發布感測值，也能透過 MQTT 訂閱主題接收 JSON 指令控制綠、黃、紅三顆 LED。顯示介面採用 Adafruit GFX 繪圖，包含總覽頁、溫度頁、濕度頁與亮度頁；後三頁以彩色圓點 Gauge 與最近 10 分鐘折線圖呈現趨勢，GPIO0 按鍵可循環換頁。")
add_heading(doc, "一、專案目標", 1)
add_bullets(doc, [
    "整合溫度、濕度、亮度三類環境資料，並在小尺寸 TFT 上清楚呈現。",
    "以 MQTT 建立雙向物聯網通訊：上傳感測 JSON，下載控制 JSON。",
    "以局部資料更新與歷史緩衝，讓畫面保持穩定且能觀察趨勢。",
    "以 NTP 同步台灣時區（UTC+8），在畫面下方顯示日期與時間。",
    "保留可擴充的分頁架構，方便加入其他感測器或控制功能。",
])
add_heading(doc, "二、系統架構", 1)
add_para(doc, "ESP32 負責感測、顯示、網路通訊與輸出控制。DHT11 及光敏元件提供資料；ILI9225 顯示器負責本地人機介面；MQTT broker 提供資料交換；三色 LED 接收遠端控制結果。NTP 則提供穩定的時間來源。")
add_picture(doc, ARCH, 6.25, "圖 2　硬體、顯示、雲端通訊與輸出控制的關係")
add_heading(doc, "三、程式流程", 1)
add_picture(doc, FLOW, 3.55, "圖 3　系統啟動、循環讀值、顯示、發布與訂閱控制流程")
add_para(doc, "系統啟動後依序連線 Wi-Fi、同步 NTP、連線 MQTT，接著讀取感測器並繪製目前頁面。主迴圈持續執行 MQTT callback、按鍵換頁及連線檢查；每 10 秒讀取一次資料、保存一筆歷史值、更新畫面並發布 JSON。歷史緩衝最多保存 60 筆，對應約 10 分鐘資料。")
doc.add_page_break()

add_heading(doc, "四、硬體與腳位配置", 1)
add_table(doc, ["模組", "功能", "ESP32 腳位 / 設定"], [
    ["ILI9225 TFT", "SPI 時脈", "GPIO18（SCK）"],
    ["ILI9225 TFT", "SPI 資料輸入 SDI/MOSI", "GPIO23（MOSI）；無 SDO/MISO"],
    ["ILI9225 TFT", "片選 CS", "GPIO27"],
    ["ILI9225 TFT", "資料/命令 DC（RS）", "GPIO26"],
    ["ILI9225 TFT", "重置 RST", "GPIO25"],
    ["DHT11", "溫度、濕度", "GPIO14"],
    ["光敏元件", "ADC 亮度，換算 0–100%", "GPIO33"],
    ["頁面按鍵", "按下換頁，內部上拉", "GPIO0 → GND"],
    ["LED", "綠 / 黃 / 紅控制", "GPIO15 / GPIO2 / GPIO4"],
], [1.35, 2.25, 2.8])
add_heading(doc, "五、軟體與通訊設定", 1)
add_table(doc, ["項目", "設定"], [
    ["顯示繪圖", "Adafruit GFX Library，自製 ILI9225 SPI 顯示類別"],
    ["Wi-Fi", "SSID：a；密碼：12345678"],
    ["MQTT broker", "mqttgo.io：1883"],
    ["發布主題", "eric/class302/data"],
    ["訂閱主題", "eric/class302/ctrl"],
    ["時間同步", "pool.ntp.org / time.nist.gov；UTC+8"],
    ["主要函式庫", "PubSubClient、ArduinoJson、SimpleDHT、Adafruit GFX"],
], [1.65, 4.75])
add_heading(doc, "六、MQTT JSON 設計", 1)
add_para(doc, "每 10 秒發布一筆目前測量值，格式固定為整數，便於儀表板或其他程式解析：")
add_para(doc, '{"temp":25,"humi":65,"light":95}', size=11, bold=True, color="17365D", before=2, after=6)
add_para(doc, "控制端將 JSON 發布到 eric/class302/ctrl，ESP32 由 mqttCallback 即時解析，避免等待下一次感測週期：")
add_para(doc, '{"gled":"on"}　　{"yled":"off"}　　{"rled":"on"}', size=11, bold=True, color="C00000", before=2, after=6)
add_para(doc, "其中 gled、yled、rled 分別對應綠燈 GPIO15、黃燈 GPIO2、紅燈 GPIO4；值為 on 時輸出 HIGH，off 時輸出 LOW。")
doc.add_page_break()

add_heading(doc, "七、TFT 四頁使用者介面", 1)
add_table(doc, ["頁面", "內容", "更新方式"], [
    ["第 1 頁：總覽", "Wi-Fi/MQTT 狀態、溫度、濕度、亮度圖示與數值、台灣日期時間", "背景 setup/換頁時繪製；數值與狀態局部更新"],
    ["第 2 頁：溫度", "10–40°C 半圓 Gauge（綠/黃/紅）與最近 10 分鐘折線圖", "每 10 秒更新"],
    ["第 3 頁：濕度", "0–100% 半圓 Gauge（綠/黃/紅）與最近 10 分鐘折線圖", "每 10 秒更新"],
    ["第 4 頁：亮度", "0–100% 半圓 Gauge（藍/黃/綠）與最近 10 分鐘折線圖", "每 10 秒更新"],
], [1.35, 3.5, 1.55])
add_para(doc, "Gauge 採用多個圓形色塊排列成半圓，再以白色內圓形成儀表盤中心；指針、刻度與數值均由 Adafruit GFX API 繪製。按下 GPIO0 按鍵即可在四頁間循環。")
add_heading(doc, "八、測試結果與成果照片", 1)
add_para(doc, "實機測試可看到 ILI9225 顯示器正常顯示狀態列、Gauge、折線圖與感測數值；照片也記錄了 ESP32、DHT11、光敏元件、LED 與 TFT 的麵包板接線成果。")
add_para(doc, "除了本地 TFT 顯示，專案歷程也保留手機端的物聯網成果畫面：手機可查看雲端圖表或接收環境異常通知。這些畫面屬於前期資料上傳/通知測試，但與目前版本的 MQTT 雙向通訊方向一致，可作為行動端延伸應用的驗證。")

photos = [
    ("WIN_20260810_13_16_18_Pro.jpg", "圖 4　0810 舊版 ESP32、DHT11、光敏電阻與 OLED 接線成果"),
    ("WIN_20260811_11_53_23_Pro.jpg", "圖 5　舊版環境監測實機測試"),
    ("WIN_20260914_11_32_42_Pro.jpg", "圖 6　0914 最新版實機接線與環境監測裝置"),
    ("WIN_20260914_13_14_28_Pro.jpg", "圖 7　溫度頁：Gauge 與最近 10 分鐘折線圖"),
    ("WIN_20260914_15_11_42_Pro.jpg", "圖 8　亮度頁：彩色 Gauge 與趨勢圖"),
    ("WIN_20260914_15_11_57_Pro.jpg", "圖 9　濕度頁：Gauge 與趨勢圖"),
    ("WIN_20260914_15_12_04_Pro.jpg", "圖 10　另一組頁面切換與接線成果"),
]
for name, caption in photos:
    add_picture(doc, PHOTO / name, 5.8, caption)
add_picture(doc, PHOTO / "Screenshot_2026-08-10-15-33-55-657_jp.naver.line.android.jpg", 3.1, "圖 11　手機端環境通知畫面（歷程成果）")
add_picture(doc, PHOTO / "螢幕擷取畫面 2026-08-10 095117.png", 6.0, "圖 12　手機/瀏覽器端溫濕度與亮度圖表畫面（歷程成果）")

doc.add_page_break()
add_heading(doc, "九、測試重點", 1)
add_bullets(doc, [
    "顯示器：確認 ILI9225 176×220 解析度、SPI 腳位及無 SDO 接法可穩定繪圖。",
    "感測器：每 10 秒讀取 DHT11 與 GPIO33 ADC，並將亮度映射至 0–100%。",
    "網路：畫面狀態列以 WIFI: O/X 與 MQTT: O/X 表示連線狀態，失聯時自動重連。",
    "控制：MQTT callback 解析 ArduinoJson，收到指令後直接切換三色 LED。",
    "趨勢：以 60 筆環形緩衝保存資料，畫出約 10 分鐘的變化曲線。",
    "時間：透過 NTP 取得 UTC+8，格式為 MM/DD 星期 HH:MM，每 10 秒刷新。",
])
add_heading(doc, "十、結論與後續改進", 1)
add_para(doc, "本版本已將感測、顯示、網路傳輸與遠端控制整合在同一個 ESP32 節點，並以四頁 TFT 介面兼顧即時數值與短期趨勢。相較於早期僅有單頁顯示或單向上傳的版本，目前系統更接近完整的物聯網終端。")
add_bullets(doc, [
    "可加入 MQTT TLS 或帳號密碼，提升公開 broker 上的通訊安全。",
    "可將歷史資料改存入 SD 卡或雲端資料庫，保留超過 10 分鐘的長期趨勢。",
    "可增加感測值異常門檻與蜂鳴器/LED 警報規則。",
    "可加入按鍵防彈跳提示、背光控制及更精細的中文字型支援。",
])
add_heading(doc, "附錄：專案檔案與版本", 1)
add_table(doc, ["項目", "內容"], [
    ["主要程式", "35_ili_mqtt_ctrl_page/35_ili_mqtt_ctrl_page.ino"],
    ["報告來源版本", "0810物聯網成果報告.docx（本檔已依 9/14 新進度改寫）"],
    ["顯示解析度", "ILI9225：176 × 220"],
    ["資料週期", "10 秒取樣、發布與畫面更新；60 筆歷史資料約 10 分鐘"],
])

doc.core_properties.title = "0914物聯網成果報告"
doc.core_properties.subject = "ESP32 ILI9225 MQTT 環境監測與控制"
doc.core_properties.author = ""
doc.core_properties.comments = "UTF-8 Python source; Microsoft JhengHei document font"
doc.save(OUT)
print(OUT)
