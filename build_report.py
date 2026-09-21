# -*- coding: utf-8 -*-
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\esp32實習")
PHOTO = ROOT / "成果照片"
OUT = ROOT / "0810物聯網成果報告_imagegen.docx"
FONT = r"C:\Windows\Fonts\msjh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msjhbd.ttc"


def fnt(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def wrap(draw, text, font, max_width):
    words = list(text) if any("\u4e00" <= c <= "\u9fff" for c in text) else text.split()
    lines, line = [], ""
    for word in words:
        candidate = line + word if words is not text.split() else (line + " " + word).strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width or not line:
            line = candidate
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def centered(draw, box, text, font, fill=(30, 45, 65)):
    x1, y1, x2, y2 = box
    bb = draw.textbbox((0, 0), text, font=font)
    draw.text(((x1 + x2 - bb[2]) / 2, (y1 + y2 - bb[3]) / 2 - bb[1]), text, font=font, fill=fill)


def box(draw, xy, title, subtitle, fill=(235, 244, 250), outline=(45, 110, 150)):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=4)
    centered(draw, (x1 + 10, y1 + 18, x2 - 10, y1 + 65), title, fnt(30, True), outline)
    for i, line in enumerate(wrap(draw, subtitle, fnt(22), x2 - x1 - 30)[:2]):
        centered(draw, (x1 + 12, y1 + 78 + i * 29, x2 - 12, y1 + 105 + i * 29), line, fnt(22), (50, 65, 80))


def arrow(draw, start, end, fill=(65, 105, 130), width=6):
    draw.line([start, end], fill=fill, width=width)
    import math
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    p1 = (end[0] - size * math.cos(ang - 0.45), end[1] - size * math.sin(ang - 0.45))
    p2 = (end[0] - size * math.cos(ang + 0.45), end[1] - size * math.sin(ang + 0.45))
    draw.polygon([end, p1, p2], fill=fill)


def make_architecture(path):
    im = Image.new("RGB", (1800, 1000), (250, 252, 253))
    d = ImageDraw.Draw(im)
    d.text((70, 35), "ESP32 物聯網環境監測系統｜系統架構圖", font=fnt(42, True), fill=(24, 65, 90))
    box(d, (70, 185, 440, 395), "感測端", "DHT11\n溫度、濕度", (231, 247, 240), (42, 135, 90))
    box(d, (70, 530, 440, 740), "感測端", "光敏電阻\n亮度 0–100%", (255, 246, 220), (194, 130, 32))
    box(d, (640, 300, 1110, 620), "ESP32 Dev Module", "GPIO14／GPIO33\n資料判斷、計時與網路傳輸", (225, 239, 250), (35, 105, 160))
    box(d, (1360, 130, 1730, 320), "OLED 顯示器", "I²C：SDA 21、SCL 22\n即時數值與通知狀態", (242, 234, 252), (115, 75, 155))
    box(d, (1360, 405, 1730, 595), "Google Sheets", "Google Apps Script\n每 10 秒寫入資料", (231, 246, 238), (45, 130, 85))
    box(d, (1360, 680, 1730, 870), "LINE 使用者", "異常推播通知\n溫度 > 28 或濕度 > 70", (255, 235, 235), (180, 65, 65))
    arrow(d, (440, 290), (640, 405)); arrow(d, (440, 635), (640, 515))
    arrow(d, (1110, 365), (1360, 225)); arrow(d, (1110, 460), (1360, 500)); arrow(d, (1110, 555), (1360, 775))
    d.text((1150, 305), "I²C", font=fnt(24, True), fill=(90, 80, 130))
    d.text((1145, 430), "Wi‑Fi / HTTPS", font=fnt(24, True), fill=(45, 110, 90))
    d.text((1140, 650), "異常判斷", font=fnt(24, True), fill=(160, 65, 65))
    im.save(path)


def make_flow(path):
    im = Image.new("RGB", (1500, 1700), (250, 252, 253))
    d = ImageDraw.Draw(im)
    d.text((70, 35), "ESP32 物聯網環境監測系統｜資料流程圖", font=fnt(42, True), fill=(24, 65, 90))
    nodes = [
        ((500, 120, 1000, 220), "開始／初始化", "啟動 OLED、ADC 與 Wi‑Fi"),
        ((500, 300, 1000, 410), "讀取感測資料", "DHT11：溫度、濕度；GPIO33：亮度"),
        ((500, 490, 1000, 600), "更新 OLED", "顯示三項數值與圖示"),
        ((500, 680, 1000, 790), "每 10 秒？", "是：上傳 Google Sheets；否：繼續"),
        ((500, 870, 1000, 980), "異常判斷", "溫度 > 28 或濕度 > 70？"),
        ((120, 1100, 620, 1220), "LINE 通知", "OLED 顯示傳送狀態，30 秒最多一次"),
        ((880, 1100, 1380, 1220), "維持監測", "返回讀取感測資料"),
        ((500, 1330, 1000, 1440), "等待 1 秒", "持續循環執行"),
    ]
    for i, (xy, title, sub) in enumerate(nodes):
        fill = (225, 239, 250) if i < 5 else ((255, 235, 235) if i == 5 else (231, 246, 238))
        box(d, xy, title, sub, fill)
    arrow(d, (750, 220), (750, 300)); arrow(d, (750, 410), (750, 490)); arrow(d, (750, 600), (750, 680)); arrow(d, (750, 790), (750, 870))
    arrow(d, (650, 925), (620, 1160), fill=(180, 65, 65)); arrow(d, (850, 925), (880, 1160), fill=(45, 130, 85))
    arrow(d, (370, 1220), (650, 1380), fill=(180, 65, 65)); arrow(d, (1130, 1220), (850, 1380), fill=(45, 130, 85)); arrow(d, (750, 1440), (750, 300))
    d.text((520, 810), "每 10 秒上傳", font=fnt(24, True), fill=(45, 110, 90))
    d.text((535, 1015), "是", font=fnt(26, True), fill=(180, 65, 65)); d.text((850, 1015), "否", font=fnt(26, True), fill=(45, 130, 85))
    im.save(path)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft JhengHei"
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r.font.size = Pt(10)
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(90, 90, 90)


def add_picture(doc, path, width=6.2, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        add_caption(doc, caption)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)
    paragraph.add_run(" 頁")


def main():
    # 架構圖與流程圖由 imagegen 產生，保留原始 PNG 並嵌入報告。
    arch = ROOT / "架構圖_imagegen.png"
    flow = ROOT / "流程圖_imagegen.png"

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
    normal = doc.styles['Normal']
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    normal.font.size = Pt(11)
    for name, size, color in [('Title', 28, (24, 65, 90)), ('Heading 1', 18, (24, 90, 120)), ('Heading 2', 14, (45, 105, 130))]:
        st = doc.styles[name]
        st.font.name = "Microsoft JhengHei"
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        st.font.size = Pt(size); st.font.color.rgb = RGBColor(*color)

    header = sec.header.paragraphs[0]
    header.text = "0810 物聯網成果報告｜ESP32 環境監測與雲端通知"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("0810\n").bold = True
    p.runs[0].font.size = Pt(18); p.runs[0].font.color.rgb = RGBColor(45, 130, 160)
    r = p.add_run("物聯網成果報告"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = RGBColor(24, 65, 90)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p2.add_run("ESP32 環境監測、Google Sheets 紀錄與 LINE 異常通知")
    rr.font.size = Pt(15); rr.font.color.rgb = RGBColor(80, 90, 100)
    doc.add_paragraph("")
    add_picture(doc, PHOTO / "IMG_20260810_090632.jpg", 4.9, "圖 1　ESP32、DHT11、光敏電阻與 OLED 實體接線成果")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("專案版本：25_dht_line\n完成日期：2026 年 8 月 10 日").font.size = Pt(11)
    doc.add_page_break()

    doc.add_heading("摘要", level=1)
    doc.add_paragraph("本專案以 ESP32 Dev Module 為核心，整合 DHT11 溫濕度感測器、光敏電阻與 I²C OLED 顯示器，完成環境資料的即時量測、視覺化呈現、Google Sheets 雲端紀錄與 LINE 異常通知。系統每秒更新感測值，每 10 秒將溫度、濕度及亮度資料寫入 Google Sheets；當溫度超過 28°C 或濕度超過 70% 時，透過 LINE Messaging API 通知指定使用者，且異常期間每 30 秒最多發送一次。")

    doc.add_heading("一、專案目標", level=1)
    for t in [
        "學習 ESP32 GPIO、ADC、I²C 與 Wi‑Fi 網路功能。",
        "完成溫度、濕度與環境亮度的即時監測。",
        "以 OLED 提供現場即時資訊與網路傳輸狀態。",
        "將感測資料寫入 Google Sheets，方便保存與製作圖表。",
        "建立環境異常判斷，透過 LINE 即時通知使用者。",
    ]:
        doc.add_paragraph(t, style='List Bullet')

    doc.add_heading("二、系統架構", level=1)
    doc.add_paragraph("系統由感測端、ESP32 控制端、OLED 顯示端及兩個雲端服務組成。ESP32 負責讀取資料、判斷異常、更新畫面並透過 Wi‑Fi 傳送資料。")
    add_picture(doc, arch, 6.45, "圖 2　系統架構圖")

    doc.add_heading("三、硬體與接線", level=1)
    table = doc.add_table(rows=1, cols=4); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
    headers = ["元件", "用途", "ESP32 腳位", "備註"]
    for c, h in zip(table.rows[0].cells, headers): set_cell_text(c, h, True, (255,255,255)); set_cell_shading(c, '287D9E')
    rows = [
        ("ESP32 Dev Module", "主控制器、Wi‑Fi", "—", "Arduino ESP32 core 3.3.10"),
        ("DHT11", "溫度、濕度", "GPIO14", "每秒讀取一次"),
        ("光敏電阻", "環境亮度", "GPIO33 / ADC", "換算為 0–100%"),
        ("OLED SSD1306", "顯示資料與狀態", "SDA 21、SCL 22", "I²C 128×64"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for c, v in zip(cells, row): set_cell_text(c, v)

    doc.add_heading("四、軟體流程", level=1)
    doc.add_paragraph("開機後系統先初始化感測器、OLED 與 Wi‑Fi，並在 OLED 顯示連線進度。進入主迴圈後，每秒讀取感測器並更新畫面；雲端紀錄與 LINE 通知各自使用獨立計時器，避免阻塞主要監測流程。")
    add_picture(doc, flow, 6.45, "圖 3　資料流程圖")

    doc.add_heading("五、核心功能與程式設計", level=1)
    features = [
        ("即時監測", "DHT11 讀取溫度與濕度，GPIO33 ADC 讀取光敏電阻，亮度以 map() 轉換成 0–100%。"),
        ("OLED 顯示", "上方顯示溫度，下方左右分別顯示濕度與亮度，並使用溫度計、水滴、太陽圖示。"),
        ("Google Sheets", "以 HTTPS 連線 Google Apps Script，資料格式為溫度、濕度、亮度，每 10 秒寫入一筆。"),
        ("LINE 警報", "溫度 > 28°C 或濕度 > 70% 時發送警報；異常持續時以 30 秒為通知間隔。"),
        ("狀態回饋", "Wi‑Fi 連線、Google Sheets 上傳及 LINE 傳訊息期間，OLED 顯示目前處理狀態。"),
    ]
    table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
    for c, h in zip(table.rows[0].cells, ["功能", "實作內容"]): set_cell_text(c, h, True, (255,255,255)); set_cell_shading(c, '287D9E')
    for a, b in features:
        cells = table.add_row().cells; set_cell_text(cells[0], a, True); set_cell_text(cells[1], b)

    doc.add_heading("六、成果展示", level=1)
    add_picture(doc, PHOTO / "IMG_20260810_090632.jpg", 5.3, "圖 4　實體電路與 OLED 顯示成果")
    doc.add_paragraph("實體成果可看到 ESP32、DHT11、光敏電阻與 OLED 已完成接線。OLED 顯示溫度、濕度與亮度三項資訊，並以圖示與分區版面提升辨識度。")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(PHOTO / "Screenshot_2026-08-10-15-33-55-657_jp.naver.line.android.jpg"), width=Inches(2.35))
    p.add_run("    ")
    p.add_run().add_picture(str(PHOTO / "螢幕擷取畫面 2026-08-10 114134.png"), width=Inches(3.35))
    add_caption(doc, "圖 5　LINE 異常通知與 Google Sheets 資料紀錄")
    add_picture(doc, PHOTO / "螢幕擷取畫面 2026-08-10 114143.png", 6.2, "圖 6　Google Sheets 溫度、濕度與亮度趨勢圖")

    doc.add_heading("七、測試結果", level=1)
    test = doc.add_table(rows=1, cols=3); test.alignment = WD_TABLE_ALIGNMENT.CENTER; test.style = 'Table Grid'
    for c, h in zip(test.rows[0].cells, ["測試項目", "預期結果", "結果"]): set_cell_text(c, h, True, (255,255,255)); set_cell_shading(c, '287D9E')
    results = [
        ("OLED 顯示", "持續顯示溫度、濕度、亮度", "通過"),
        ("Google Sheets", "每 10 秒新增一筆資料", "通過；可看到時間戳記與三項數值"),
        ("溫度異常", "溫度 > 28°C 時通知 LINE", "已完成通知流程"),
        ("濕度異常", "濕度 > 70% 時通知 LINE", "成果截圖顯示濕度 71% 警報"),
        ("通知間隔", "異常期間每 30 秒通知一次", "程式已設定 30 秒計時器"),
    ]
    for row in results:
        cells = test.add_row().cells
        for c, v in zip(cells, row): set_cell_text(c, v)

    doc.add_heading("八、學習成果", level=1)
    doc.add_paragraph("本專案從單顆 LED、按鈕與類比輸入等基礎練習，逐步整合到感測器、OLED、Wi‑Fi、Google Sheets 與 LINE Messaging API。過程中實際練習了 GPIO 與 ADC 讀值、I²C 顯示、HTTPS 傳輸、JSON／HTTP 概念、計時器設計及異常通知邏輯，完成從硬體接線到雲端應用的完整物聯網流程。")

    doc.add_heading("九、結論與未來改進", level=1)
    doc.add_paragraph("本系統已完成環境資料量測、OLED 即時顯示、Google Sheets 紀錄與 LINE 異常通知，達成物聯網環境監測的專案目標。後續可加入資料平均與濾波、斷線自動重連、異常恢復通知、RTC 時間校正、電池供電與更安全的憑證管理。")
    doc.add_paragraph("安全性提醒：實際部署時應避免將 Wi‑Fi 密碼、Google Apps Script 識別資訊與 LINE Channel Access Token 放在公開程式碼或報告中，並建議定期更新已使用過的 Token。")

    doc.add_heading("附錄：專案檔案", level=1)
    doc.add_paragraph("本次結案版本：25_dht_line/25_dht_line.ino\n主要硬體：ESP32 Dev Module、DHT11、光敏電阻、SSD1306 OLED\n主要函式庫：WiFi、WiFiClientSecure、SimpleDHT、U8g2lib")

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
